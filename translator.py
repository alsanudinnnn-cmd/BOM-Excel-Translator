"""Excel translation engine backed by Azure Translator."""

from __future__ import annotations

import os
import re
import sys
import tempfile
import uuid
from pathlib import Path
from typing import Iterator

import pandas as pd
import requests
from dotenv import load_dotenv
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

from usage_tracker import add_usage, count_azure_characters

APP_DIR = Path(sys.executable).resolve().parent if getattr(sys, "frozen", False) else Path(__file__).resolve().parent
load_dotenv(APP_DIR / ".env")

ALLOWED_EXTENSIONS = {".xls", ".xlsx"}
MAX_TEXTS_PER_REQUEST = 100
MAX_CHARS_PER_REQUEST = 45_000
FORMULA_PATTERN = re.compile(r"^=")


def is_translatable(value: object) -> bool:
    return isinstance(value, str) and bool(value.strip()) and not FORMULA_PATTERN.match(value)


def batches(texts: list[str]) -> Iterator[list[str]]:
    batch: list[str] = []
    char_count = 0
    for item in texts:
        if len(item) > MAX_CHARS_PER_REQUEST:
            raise ValueError("A cell contains more than 45,000 characters and cannot be translated.")
        if batch and (len(batch) >= MAX_TEXTS_PER_REQUEST or char_count + len(item) > MAX_CHARS_PER_REQUEST):
            yield batch
            batch, char_count = [], 0
        batch.append(item)
        char_count += len(item)
    if batch:
        yield batch


def translate_texts(texts: list[str], target_language: str, source_language: str | None) -> dict[str, str]:
    key = os.getenv("AZURE_TRANSLATOR_KEY")
    if not key:
        raise RuntimeError("AZURE_TRANSLATOR_KEY is missing. Add it to the .env file beside the app.")

    endpoint = os.getenv("AZURE_TRANSLATOR_ENDPOINT", "https://api.cognitive.microsofttranslator.com").rstrip("/")
    region = os.getenv("AZURE_TRANSLATOR_REGION", "").strip()
    headers = {"Ocp-Apim-Subscription-Key": key, "Content-Type": "application/json"}
    if region:
        headers["Ocp-Apim-Subscription-Region"] = region

    translated: dict[str, str] = {}
    for group in batches(texts):
        params = [("api-version", "3.0"), ("to", target_language)]
        if source_language:
            params.append(("from", source_language))
        response = None
        try:
            response = requests.post(
                f"{endpoint}/translate",
                params=params,
                headers=headers,
                json=[{"Text": text} for text in group],
                timeout=60,
            )
            response.raise_for_status()
        except requests.RequestException as error:
            detail = response.text if response is not None else str(error)
            raise RuntimeError(f"Azure Translator request failed: {detail}") from error

        # Azure has accepted this batch, so its source characters are billable.
        add_usage(group)
        payload = response.json()
        for original, item in zip(group, payload, strict=True):
            translated[original] = item["translations"][0]["text"]
    return translated


def convert_xls_to_xlsx(source: Path) -> Path:
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as temporary:
        destination = Path(temporary.name)
    with pd.ExcelFile(source, engine="xlrd") as excel_file:
        with pd.ExcelWriter(destination, engine="openpyxl") as writer:
            for sheet_name in excel_file.sheet_names:
                frame = pd.read_excel(excel_file, sheet_name=sheet_name, header=None, dtype=object)
                frame.to_excel(writer, sheet_name=sheet_name, index=False, header=False)
    return destination


def _collect_translatable_content(workbook) -> tuple[list, list[str]]:
    cells = []
    unique_texts: dict[str, None] = {}
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if not isinstance(cell, MergedCell) and is_translatable(cell.value):
                    cells.append(cell)
                    unique_texts.setdefault(cell.value, None)
    return cells, list(unique_texts)


def estimate_workbook_characters(input_file: Path) -> int:
    """Estimate the source characters this app will submit to Azure."""
    converted_file: Path | None = None
    try:
        if input_file.suffix.lower() == ".xls":
            converted_file = convert_xls_to_xlsx(input_file)
        workbook = load_workbook(converted_file or input_file, data_only=False, read_only=True)
        _, unique_texts = _collect_translatable_content(workbook)
        workbook.close()
        return sum(count_azure_characters(text) for text in unique_texts)
    finally:
        if converted_file is not None:
            converted_file.unlink(missing_ok=True)


def translate_workbook(
    input_file: Path,
    target_language: str,
    source_language: str | None,
    output_path: Path | None = None,
) -> Path:
    converted_file: Path | None = None
    try:
        if input_file.suffix.lower() == ".xls":
            converted_file = convert_xls_to_xlsx(input_file)
        workbook = load_workbook(converted_file or input_file, data_only=False)

        cells, unique_texts = _collect_translatable_content(workbook)
        translations = translate_texts(unique_texts, target_language, source_language)
        for cell in cells:
            cell.value = translations[cell.value]

        if output_path is None:
            output_dir = APP_DIR / "outputs"
            output_dir.mkdir(exist_ok=True)
            output_path = output_dir / f"{input_file.stem}_{target_language}_{uuid.uuid4().hex[:8]}.xlsx"
        workbook.save(output_path)
        return output_path
    finally:
        if converted_file is not None:
            converted_file.unlink(missing_ok=True)
